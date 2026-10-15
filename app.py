# -*- coding: utf-8 -*-
"""
Shop YoY 分析系統
=================

用法：
    python app.py                     # 本機執行（預設 http://127.0.0.1:5000）
    PORT=8080 gunicorn app:app        # 生產環境（Docker / Zeabur）

功能：
    1. 上載 Hugo_by Shop_YoY 匯出的 Excel
       （欄位：Month / Shop / Sold Qty / Net Sales / COGS / GP$ / GP%）
    2. 自動對照內置 stores-template.csv（店鋪主檔），補上 Regional / Class / Size / OM / Type 等屬性
    3. 依 Site、年份、月份（單月或多月累計）計算 YoY%，支援 Regional / OM 等進階篩選及分組
    4. 匯出 Excel 報告
"""

import csv
import gzip
import io
import json
import os
import re
import threading
import time
import uuid
import zipfile
from datetime import date, datetime

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from flask import Flask, Response, jsonify, render_template, request, send_file

try:  # 快速 Excel 讀取引擎（Rust 實作，比 openpyxl 快 10 倍以上）
    from python_calamine import CalamineWorkbook
except Exception:  # pragma: no cover - 未安裝時自動退回 openpyxl
    CalamineWorkbook = None

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
TEMPLATE_CSV = os.path.join(BASE_DIR, 'stores-template.csv')

MONTH_RE = re.compile(r'^(\d{4})-(0[1-9]|1[0-2])$')

app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = 32 * 1024 * 1024  # 上載上限 32 MB
app.config['TEMPLATES_AUTO_RELOAD'] = True

# 私隱設定（可用環境變數覆寫）：
#   APP_PASSWORD    設定後，開啟網站需輸入密碼（用戶名稱隨意，密碼 = 此值）
#   DATA_TTL_HOURS  上載資料於記憶體保留幾多小時後自動刪除（預設 8；0 = 不自動過期）
APP_PASSWORD = (os.environ.get('APP_PASSWORD') or '').strip()
try:
    DATA_TTL_HOURS = float(os.environ.get('DATA_TTL_HOURS', '8'))
except ValueError:
    DATA_TTL_HOURS = 8.0

# 已上載資料（只存記憶體，不寫入硬碟／資料庫），保留最近數份
_DATASETS = {}
_DATASET_ORDER = []
_TIMES = {}
_DATASETS_MAX = 5
_LOCK = threading.Lock()

# 上載解析任務（背景解析 + 進度查詢）：task_id -> 狀態
_TASKS = {}
_TASKS_LOCK = threading.Lock()


def _prune_tasks():
    """清走超過 1 小時嘅解析任務記錄。"""
    now = time.time()
    with _TASKS_LOCK:
        for tid in [t for t, v in _TASKS.items() if now - v.get('ts', 0) > 3600]:
            _TASKS.pop(tid, None)

GROUP_FIELDS = {
    'shop': None,
    'regional': 'regional',
    'om': 'om',
    'class1': 'class1',
    'class2': 'class2',
    'size': 'size',
    'type': 'type',
}
GROUP_LABELS = {
    'shop': '店鋪',
    'regional': 'Regional',
    'om': 'OM',
    'class1': 'Class 1',
    'class2': 'Class 2',
    'size': 'Size',
    'type': 'Type',
    # Hugo_all_YoY 明細檔的額外分析維度
    'brand': '品牌 Brand',
    'mch1': '大分類 MCH1',
    'mch2': '中分類 MCH2',
    'mc': '品類 MC',
    'nature': '客群 Nature',
    'postal': '客源 Postal',
    'buyer': '買手 Buyer',
    'sst': 'SST',
    'status': '貨品狀態 Status',
    'customer': '客戶類型 Customer',
    'promoid1': '推廣 PromoID1',
    'promoid2': '推廣 PromoID2',
}

# 精簡 SKU 店鋪（進階篩選快速選用；如日後名單有變，改呢個清單即可）
LEAN_SKU_SITES = {
    'HA19', 'HA21', 'HA33', 'HA37', 'HB30', 'HC44', 'HB86', 'HB77', 'HB91',
    'HB94', 'HC26', 'HC60', 'HC62', 'HC68', 'HC13', 'HC69', 'HB87', 'HB80',
}

# 按月矩陣的可用指標：metric -> (顯示種類, 標籤)
MATRIX_METRICS = {
    'salesYoy': ('yoy', '淨銷售 YoY%'),
    'qtyYoy': ('yoy', '銷量 YoY%'),
    'gpYoy': ('yoy', '毛利 GP$ YoY%'),
    'gpPp': ('pp', '毛利率 ±pp'),
    'salesVal': ('money', '淨銷售'),
    'qtyVal': ('qty', '銷量'),
    'gpVal': ('money', '毛利 GP$'),
    'gpPctVal': ('pct', '毛利率（值）'),
}


# ---------------------------------------------------------------------------
# 小工具
# ---------------------------------------------------------------------------

def _norm(v):
    """正規化欄位名：去空白／底線、轉小寫。"""
    return re.sub(r'[\s_]+', '', str(v or '')).lower()


def _num(v):
    """把儲存格值轉成 float；空白或非數字一律當 0。"""
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace(',', '').strip()
    if not s:
        return 0.0
    try:
        return float(s)
    except ValueError:
        return 0.0


def _month_str(v):
    """把 Month 欄位值轉成 'YYYY-MM'；不是有效月份就回傳空字串。"""
    if isinstance(v, (datetime, date)):
        return v.strftime('%Y-%m')
    s = str(v).strip()
    return s if MONTH_RE.match(s) else ''


def _period_label(months):
    """把月份清單化成易讀字串，如 '2026-01 至 2026-09'。"""
    if not months:
        return '（無）'
    nums = sorted({int(m[5:7]) for m in months})
    year = sorted(months)[0][:4]
    ranges = []
    s = p = nums[0]
    for n in nums[1:]:
        if n == p + 1:
            p = n
        else:
            ranges.append((s, p))
            s = p = n
    ranges.append((s, p))
    parts = []
    for a, b in ranges:
        if a == b:
            parts.append('%s-%02d' % (year, a))
        elif b == a + 1:
            parts.append('%s-%02d、%s-%02d' % (year, a, year, b))
        else:
            parts.append('%s-%02d 至 %s-%02d' % (year, a, year, b))
    return '、'.join(parts)


def load_template():
    """讀取內置店鋪主檔 stores-template.csv（每次上載時重新讀取）。"""
    tpl = {}
    if not os.path.exists(TEMPLATE_CSV):
        return tpl
    with open(TEMPLATE_CSV, encoding='utf-8-sig', newline='') as f:
        for r in csv.DictReader(f):
            code = (r.get('Site') or '').strip()
            if not code:
                continue
            tpl[code] = {
                'name': (r.get('Shop') or '').strip(),
                'regional': (r.get('Regional') or '').strip(),
                'class1': (r.get('Class 1') or '').strip(),
                'class2': (r.get('Class 2') or '').strip(),
                'size': (r.get('Size') or '').strip(),
                'om': (r.get('OM') or '').strip(),
                'type': (r.get('Type') or '').strip(),
                'focus': bool((r.get('需特別關注') or '').strip()),
            }
    return tpl


# ---------------------------------------------------------------------------
# Excel 解析
# ---------------------------------------------------------------------------

ALL_DIMS = ['brand', 'mch1', 'mch2', 'mc', 'nature', 'postal', 'buyer', 'sst',
            'status', 'customer', 'promoid1', 'promoid2']


def _count_rows_quick(data, sheet_name):
    """快速估算工作表總行數（直接數 XML 嘅 row 標籤，支援任何命名空間前綴），失敗回傳 0。"""
    try:
        row_re = re.compile(rb'<[A-Za-z0-9]*:?row[ >/]')
        zf = zipfile.ZipFile(io.BytesIO(data))
        sheet_path = None
        try:
            from xml.etree import ElementTree as ET
            rid = None
            wbxml = ET.fromstring(zf.read('xl/workbook.xml'))
            for el in wbxml.iter():
                if el.tag.endswith('sheet') and el.get('name') == sheet_name:
                    rid = el.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
                    break
            if rid:
                rels = ET.fromstring(zf.read('xl/_rels/workbook.xml.rels'))
                for rel in rels.iter():
                    if rel.tag.endswith('Relationship') and rel.get('Id') == rid:
                        tgt = rel.get('Target') or ''
                        if tgt.startswith('/'):
                            sheet_path = tgt.lstrip('/')
                        elif tgt.startswith('xl/'):
                            sheet_path = tgt
                        else:
                            sheet_path = 'xl/' + tgt
                        break
        except Exception:
            sheet_path = None
        if not sheet_path or sheet_path not in zf.namelist():
            cands = [n for n in zf.namelist() if n.startswith('xl/worksheets/') and n.endswith('.xml')]
            sheet_path = cands[0] if cands else None
        if not sheet_path:
            return 0
        n = 0
        tail = b''
        with zf.open(sheet_path) as f:
            while True:
                chunk = f.read(1 << 20)
                if not chunk:
                    break
                buf = tail + chunk
                n += len(row_re.findall(buf))
                tail = buf[-16:]
        return n
    except Exception:
        return 0


def _header_colmap(row):
    """如果呢行係表頭（含 Month / Shop / Net Sales）就回傳欄位對照，唔係就 None。"""
    cells = {}
    for ci, cv in enumerate(row or ()):
        if cv is None or cv == '':
            continue
        cells[_norm(cv)] = ci
    if 'month' in cells and 'shop' in cells and 'netsales' in cells:
        return cells
    return None


def _scan_sheets_calamine(wb):
    for name in wb.sheet_names:
        ws = wb.get_sheet_by_name(name)
        it = ws.iter_rows()
        for ri in range(1, 11):
            row = next(it, None)
            if row is None:
                break
            cm = _header_colmap(row)
            if cm:
                try:
                    total = int(ws.height)  # calamine 直接提供總行數（即時）
                except Exception:
                    total = 0
                return name, ri, cm, it, None, total
    return None


def _scan_sheets_openpyxl(wb):
    for sh in wb.worksheets:
        if getattr(sh, 'sheet_state', 'visible') != 'visible':
            continue
        it = sh.iter_rows(values_only=True)
        for ri in range(1, 11):
            row = next(it, None)
            if row is None:
                break
            cm = _header_colmap(row)
            if cm:
                return sh.title, ri, cm, it, wb.close, 0
    return None


def _find_header_reader(data):
    """開啟 Excel 並搵表頭（每張工作表頭 10 行內）。

    回傳 (sheet_title, header_index, colmap, row_iter, close_fn, total_hint)；
    開得到但搵唔到表頭就回傳 None（total_hint > 0 表示唔使再自行數行）。
    優先用 python-calamine（快），失敗或唔支援就退回 openpyxl。
    """
    if CalamineWorkbook is not None:
        try:
            wb = CalamineWorkbook.from_filelike(io.BytesIO(data))
            found = _scan_sheets_calamine(wb)
            if found:
                return found
        except Exception:
            pass
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise ValueError('無法讀取檔案，請確認為有效的 Excel (.xlsx) 檔案。')
    found = _scan_sheets_openpyxl(wb)
    if not found:
        wb.close()
    return found


def parse_excel(data, progress=None):
    """
    解析 Excel（bytes），回傳 dataset。自動偵測兩種格式：

    1. Hugo_by Shop_YoY（Month / Shop / …）→ kind='shop'
       rows = [月份, 店鋪, 銷量, 淨銷售, 毛利]
    2. Hugo_all_YoY（含 Brand / MCH1 / MCH2 / MC / Nature / Postal 等維度）→ kind='rich'
       rows 末端追加各分析維度索引；可用維度見 ds['dims']

    讀取優先使用 python-calamine（快 10 倍以上）；未安裝或失敗時自動退回 openpyxl。
    progress(pct, msg)：可選進度回呼（0-100）。
    """
    def cb(pct, msg=''):
        if progress:
            try:
                progress(min(100.0, max(0.0, float(pct))), msg)
            except Exception:
                pass

    cb(1, '讀取檔案…')
    cb(3, '載入數據…')
    found = _find_header_reader(data)
    if found is None:
        raise ValueError(
            '找不到資料表：檔案需包含 Month、Shop、Net Sales 欄位（Hugo_by Shop_YoY 或 Hugo_all_YoY 匯出格式）。')
    sheet_title, header_i, colmap, row_iter, close_fn, total_hint = found

    try:
        # calamine 直接提供總行數；後備路徑先自行快速數行（用於進度顯示）
        total_rows = total_hint if total_hint else _count_rows_quick(data, sheet_title)
        cb(35, '解析資料…')

        i_month = colmap['month']
        i_shop = colmap['shop']
        i_sales = colmap['netsales']
        i_qty = colmap.get('soldqty')
        i_cogs = colmap.get('cogs')
        i_gp = colmap.get('gp$')
        # 額外分析維度（Hugo_all_YoY 明細檔才有）
        dim_cols = {k: colmap.get(k) for k in ALL_DIMS}
        rich = any(v is not None for v in dim_cols.values())

        def cell(row, i):
            return row[i] if (i is not None and i < len(row)) else None

        def cell_str(row, i):
            v = cell(row, i)
            return str(v).strip() if v is not None else ''

        agg = {}            # (month, code[, 各維度值…]) -> [qty, sales, gp]
        excel_names = {}    # code -> 店名（取自 Excel 的 'CODE - 店名' 格式）
        months_set = set()
        dim_valuesets = {k: set() for k in ALL_DIMS}

        processed = 0
        for row in row_iter:
            processed += 1
            if total_rows and processed % 2000 == 0:
                cb(35 + 58 * min(1.0, processed / total_rows),
                   '解析中… %s / %s 行' % (format(processed, ','), format(total_rows, ',')))
            if row is None or len(row) <= max(i_month, i_shop):
                continue
            m = _month_str(row[i_month])
            if not m:
                continue  # 排除 Total 行、備註行等
            shop_raw = row[i_shop]
            if shop_raw is None or not str(shop_raw).strip():
                continue
            shop = str(shop_raw).strip()
            if ' - ' in shop:
                code, name = shop.split(' - ', 1)
                code, name = code.strip(), name.strip()
            else:
                code, name = shop, ''
            if code not in excel_names and name:
                excel_names[code] = name

            qty = _num(cell(row, i_qty))
            sales = _num(cell(row, i_sales))
            if i_gp is not None:
                gp = _num(cell(row, i_gp))
            else:
                gp = sales - _num(cell(row, i_cogs))

            if rich:
                dv = []
                for k in ALL_DIMS:
                    v = cell_str(row, dim_cols[k])
                    dim_valuesets[k].add(v)
                    dv.append(v)
                key = (m, code) + tuple(dv)
            else:
                key = (m, code)
            t = agg.get(key)
            if t is None:
                agg[key] = [qty, sales, gp]
            else:
                t[0] += qty
                t[1] += sales
                t[2] += gp
            months_set.add(m)

        if not agg:
            raise ValueError('找不到有效的資料列（Month 欄位需為 YYYY-MM 格式）。')

        cb(95, '對照店鋪主檔…')
        months = sorted(months_set)
        m_idx = {m: i for i, m in enumerate(months)}

        # 對照店鋪主檔
        tpl = load_template()
        sites = []
        s_idx = {}
        for code in sorted({k[1] for k in agg}):
            t = tpl.get(code)
            sites.append({
                'code': code,
                'name': (t or {}).get('name') or excel_names.get(code, '') or code,
                'regional': (t or {}).get('regional', ''),
                'class1': (t or {}).get('class1', ''),
                'class2': (t or {}).get('class2', ''),
                'size': (t or {}).get('size', ''),
                'om': (t or {}).get('om', ''),
                'type': (t or {}).get('type', ''),
                'focus': bool((t or {}).get('focus', False)),
                'inTemplate': t is not None,
                'leanSku': code in LEAN_SKU_SITES,
            })
            s_idx[code] = len(sites) - 1

        # 可分析維度：有欄位且值多過一個先至有用
        dims = {}
        if rich:
            pos = 5
            for k in ALL_DIMS:
                if dim_cols[k] is None:
                    continue
                vals = sorted(dim_valuesets[k])
                if len(vals) < 2:
                    continue
                dims[k] = {'values': vals, 'pos': pos, 'idx': {v: i for i, v in enumerate(vals)}}
                pos += 1

        cb(97, '整理數據…')
        # 預先整理維度索引（避免逐行查 dict）
        pickers = []
        if rich:
            for j, k in enumerate(ALL_DIMS):
                d = dims.get(k)
                if d is not None:
                    pickers.append((d['idx'], j))
        rows = []
        n_keys = len(agg)
        step = max(1, n_keys // 50)
        i = 0
        for key, (qty, sales, gp) in agg.items():
            i += 1
            m, code = key[0], key[1]
            r = [m_idx[m], s_idx[code], int(round(qty)), round(sales, 2), round(gp, 2)]
            if rich:
                dim_vals = key[2:]
                for idx_map, j in pickers:
                    r.append(idx_map[dim_vals[j]])
            rows.append(r)
            if n_keys > 100000 and i % step == 0:
                cb(97 + 1.5 * i / n_keys,
                   '整理數據… %s / %s' % (format(i, ','), format(n_keys, ',')))
        rows.sort(key=lambda r: (r[0], r[1]))

        years = {}
        for m in months:
            years.setdefault(m[:4], []).append(m)

        ds = {
            'kind': 'rich' if rich else 'shop',
            'months': months,
            'years': years,
            'sites': sites,
            'rows': rows,
            'latestMonth': months[-1],
            'rowCount': len(rows),
            'siteCount': len(sites),
        }
        if rich:
            ds['dims'] = {k: {'values': d['values'], 'pos': d['pos']} for k, d in dims.items()}
        cb(99, '即將完成…')
        return ds
    finally:
        if close_fn:
            close_fn()


# ---------------------------------------------------------------------------
# YoY 計算（匯出用；與前端 index.html 的邏輯一致）
# ---------------------------------------------------------------------------

def _resolve_levels(ds, opts):
    """回傳分組層級清單（最多 3 層）。支援 groupBys（多層）或 groupBy（單層，兼容舊版）。"""
    raw = opts.get('groupBys')
    if not raw:
        raw = [opts.get('groupBy') or 'shop']
    dims = ds.get('dims') or {}
    levels = []
    for g in raw:
        g = str(g or '')
        if (g in GROUP_FIELDS or g in dims) and g not in levels:
            levels.append(g)
        if len(levels) >= 3:
            break
    return levels or ['shop']


def _prepare(ds, opts):
    """建立店鋪篩選結果及分組容器。

    回傳 dict：
      levels      分組層級清單（最多 3 層）
      site_mode   True = 單層主檔屬性／店鋪（沿用預先建立組別的模式）
      groups      組別容器（site_mode 預先建立；否則由 row 逐個建立）
      site_group  site index -> group key（只限 site_mode）
      pass_set    通過篩選的 site index 集合
      pass_count  通過篩選的店鋪數
    """
    sites = ds['sites']
    levels = _resolve_levels(ds, opts)
    site_mode = len(levels) == 1 and levels[0] in GROUP_FIELDS

    sel_codes = opts.get('siteCodes')
    sel_codes = {str(c) for c in sel_codes} if sel_codes else None
    fsets = {k: {str(v) for v in (opts.get(k) or [])}
             for k in ('regional', 'om', 'class1', 'class2', 'size', 'type')}

    def site_pass(s):
        if sel_codes is not None and s['code'] not in sel_codes:
            return False
        if opts.get('templateOnly') and not s['inTemplate']:
            return False
        if opts.get('focusOnly') and not s['focus']:
            return False
        if opts.get('leanSkuOnly') and not s.get('leanSku'):
            return False
        for k, vals in fsets.items():
            if vals and (s.get(k) or '') not in vals:
                return False
        return True

    groups = {}
    site_group = {}
    pass_set = set()
    if site_mode:
        gfield = GROUP_FIELDS.get(levels[0])
        for i, s in enumerate(sites):
            if not site_pass(s):
                continue
            pass_set.add(i)
            if gfield is None:
                key = 'site:' + s['code']
                groups[key] = {
                    'label': (s['code'] + (' ' + s['name'] if s['name'] else '')).strip(),
                    'depth': 0, 'key': [key],
                    'code': s['code'], 'name': s['name'],
                    'regional': s['regional'], 'om': s['om'],
                    'class1': s['class1'], 'class2': s['class2'],
                    'size': s['size'], 'type': s['type'], 'focus': s['focus'],
                    'n': 1, 'cur': [0.0, 0.0, 0.0], 'ly': [0.0, 0.0, 0.0],
                }
            else:
                val = (s.get(gfield) or '').strip() or '（未分類）'
                key = gfield + ':' + val
                g = groups.get(key)
                if g is None:
                    g = {'label': val, 'depth': 0, 'key': [key], 'n': 0,
                         'cur': [0.0, 0.0, 0.0], 'ly': [0.0, 0.0, 0.0]}
                    groups[key] = g
                g['n'] += 1
            site_group[i] = key
    else:
        for i, s in enumerate(sites):
            if site_pass(s):
                pass_set.add(i)

    return {
        'levels': levels, 'site_mode': site_mode,
        'groups': groups, 'site_group': site_group, 'pass_set': pass_set,
        'pass_count': len(pass_set),
    }


def _make_group(ds, levels, key, si):
    """建立一個層級組別（key = 由頂層到本層嘅值 tuple）。"""
    depth = len(key) - 1
    lv = levels[depth]
    val = key[-1]
    g = {'depth': depth, 'key': list(key), 'n': 0,
         'cur': [0.0, 0.0, 0.0], 'ly': [0.0, 0.0, 0.0], '_shops': set()}
    if lv == 'shop':
        s = ds['sites'][val]
        g['label'] = (s['code'] + (' ' + s['name'] if s['name'] else '')).strip()
        g['site'] = s
    elif lv in GROUP_FIELDS:
        g['label'] = val or '（未分類）'
    else:
        g['label'] = ds['dims'][lv]['values'][val] or '（未分類）'
    return g


def _row_groups(ds, prep, groups, row, si):
    """取得／建立該 row 所屬嘅所有層級組別（由頂層至最底層，每層都要累加）。"""
    levels = prep['levels']
    sites = ds['sites']
    vals = []
    for lv in levels:
        if lv == 'shop':
            vals.append(si)
        elif lv in GROUP_FIELDS:
            vals.append((sites[si].get(lv) or '').strip() or '（未分類）')
        else:
            d = ds['dims'][lv]
            vals.append(row[d['pos']])
    out = []
    for d in range(len(vals)):
        key = tuple(vals[:d + 1])
        g = groups.get(key)
        if g is None:
            g = _make_group(ds, levels, key, si)
            groups[key] = g
        if si not in g['_shops']:
            g['_shops'].add(si)
            g['n'] += 1
        out.append(g)
    return out


def _sort_hierarchy(out):
    """把組別排成階層次序（每層按本期淨銷售 desc；用 g['cur']）。"""
    children = {}
    roots = []
    for g in out:
        if len(g['key']) == 1:
            roots.append(g)
        else:
            children.setdefault(tuple(g['key'][:-1]), []).append(g)
    roots.sort(key=lambda g: (-g['cur'][1], g['label']))
    result = []

    def emit(g):
        result.append(g)
        kids = children.get(tuple(g['key']))
        if kids:
            kids.sort(key=lambda x: (-x['cur'][1], x['label']))
            for k in kids:
                emit(k)

    for r in roots:
        emit(r)
    return result


def compute_report(ds, opts):
    """累計模式：所選月份合併，與去年同期比較（支援多層分組）。"""
    year = int(opts.get('year'))
    lyear = year - 1
    months = {str(m) for m in (opts.get('months') or [])}
    month_nums = {m[5:7] for m in months}
    months_list = ds['months']

    prep = _prepare(ds, opts)
    groups = prep['groups']
    site_group = prep['site_group']
    pass_set = prep['pass_set']
    site_mode = prep['site_mode']

    ystr, lystr = str(year), str(lyear)
    for row in ds['rows']:
        si = row[1]
        if si not in pass_set:
            continue
        m = months_list[row[0]]
        if m[:4] == ystr and m in months:
            target = 'cur'
        elif m[:4] == lystr and m[5:7] in month_nums:
            target = 'ly'
        else:
            continue
        if site_mode:
            gs = (groups[site_group[si]],)
        else:
            gs = _row_groups(ds, prep, groups, row, si)
        for g in gs:
            t = g[target]
            t[0] += row[2]
            t[1] += row[3]
            t[2] += row[4]

    hide_empty = opts.get('hideEmpty', True)
    out = []
    for g in groups.values():
        if hide_empty and not any(g['cur']) and not any(g['ly']):
            continue
        out.append(g)
    if site_mode:
        out.sort(key=lambda g: (-g['cur'][1], g['label']))
    else:
        out = _sort_hierarchy(out)

    totals = {'cur': [0.0, 0.0, 0.0], 'ly': [0.0, 0.0, 0.0]}
    for g in out:
        if g['depth'] != 0:  # 只加總頂層，避免重複計算
            continue
        for j in range(3):
            totals['cur'][j] += g['cur'][j]
            totals['ly'][j] += g['ly'][j]

    return {'groups': out, 'totals': totals, 'siteCount': prep['pass_count']}


def compute_matrix(ds, opts):
    """按月矩陣：每個組別 × 每個所選月份（支援多層分組）。"""
    year = int(opts.get('year'))
    lyear = year - 1
    months = sorted({str(m) for m in (opts.get('months') or [])})
    ly_of = {('%d-%s' % (lyear, m[5:7])): m for m in months}
    months_list = ds['months']

    prep = _prepare(ds, opts)
    groups = prep['groups']
    site_group = prep['site_group']
    pass_set = prep['pass_set']
    site_mode = prep['site_mode']

    def ensure_cells(g):
        if 'cells' not in g:
            g['cells'] = {m: [0.0, 0.0, 0.0] for m in months}
            g['cellsLy'] = {m: [0.0, 0.0, 0.0] for m in months}
            g['aggCur'] = [0.0, 0.0, 0.0]
            g['aggLy'] = [0.0, 0.0, 0.0]

    for row in ds['rows']:
        si = row[1]
        if si not in pass_set:
            continue
        m = months_list[row[0]]
        if m in months:
            tname, cm = 'cells', m
        elif m in ly_of:
            tname, cm = 'cellsLy', ly_of[m]
        else:
            continue
        if site_mode:
            gs = (groups[site_group[si]],)
        else:
            gs = _row_groups(ds, prep, groups, row, si)
        for g in gs:
            ensure_cells(g)
            cell = g[tname][cm]
            cell[0] += row[2]
            cell[1] += row[3]
            cell[2] += row[4]

    for g in groups.values():
        ensure_cells(g)
    for g in groups.values():
        for m in months:
            for j in range(3):
                g['aggCur'][j] += g['cells'][m][j]
                g['aggLy'][j] += g['cellsLy'][m][j]

    hide_empty = opts.get('hideEmpty', True)
    out = []
    for g in groups.values():
        if hide_empty and not any(g['aggCur']) and not any(g['aggLy']):
            continue
        out.append(g)
    if site_mode:
        out.sort(key=lambda g: (-g['aggCur'][1], g['label']))
    else:
        for g in out:  # _sort_hierarchy 用 cur 排序
            g['cur'] = g['aggCur']
        out = _sort_hierarchy(out)

    totals = {
        'cells': {m: [0.0, 0.0, 0.0] for m in months},
        'cellsLy': {m: [0.0, 0.0, 0.0] for m in months},
        'aggCur': [0.0, 0.0, 0.0],
        'aggLy': [0.0, 0.0, 0.0],
    }
    for g in out:
        if g['depth'] != 0:  # 只加總頂層，避免重複計算
            continue
        for m in months:
            for j in range(3):
                totals['cells'][m][j] += g['cells'][m][j]
                totals['cellsLy'][m][j] += g['cellsLy'][m][j]
        for j in range(3):
            totals['aggCur'][j] += g['aggCur'][j]
            totals['aggLy'][j] += g['aggLy'][j]

    return {'groups': out, 'months': months, 'totals': totals, 'siteCount': prep['pass_count']}


def _matrix_value(cur, ly, metric):
    """按月矩陣的單格數值。cur / ly = [qty, sales, gp]；無法計算回傳 None。"""
    if metric == 'qtyYoy':
        return ((cur[0] - ly[0]) / ly[0]) if ly[0] > 0 else None
    if metric == 'gpYoy':
        return ((cur[2] - ly[2]) / ly[2]) if ly[2] > 0 else None
    if metric == 'gpPp':
        pc = (cur[2] / cur[1]) if cur[1] else None
        pl = (ly[2] / ly[1]) if ly[1] else None
        return (pc - pl) if (pc is not None and pl is not None) else None
    if metric == 'salesVal':
        return cur[1]
    if metric == 'qtyVal':
        return cur[0]
    if metric == 'gpVal':
        return cur[2]
    if metric == 'gpPctVal':
        return (cur[2] / cur[1]) if cur[1] else None
    # 預設：淨銷售 YoY%
    return ((cur[1] - ly[1]) / ly[1]) if ly[1] > 0 else None


def _filter_summary(opts):
    parts = ['Site：%s' % ('全部' if not opts.get('siteCodes') else '%d 間' % len(opts['siteCodes']))]
    for label, key in (('Regional', 'regional'), ('OM', 'om'), ('Class 1', 'class1'),
                       ('Class 2', 'class2'), ('Size', 'size'), ('Type', 'type')):
        vals = opts.get(key) or []
        if vals:
            parts.append('%s：%s' % (label, '、'.join(str(v) for v in vals)))
    if opts.get('focusOnly'):
        parts.append('只顯示需特別關注店鋪')
    if opts.get('leanSkuOnly'):
        parts.append('只顯示精簡SKU店鋪')
    if opts.get('templateOnly'):
        parts.append('只顯示 stores-template 內店鋪')
    return '；'.join(parts)


def _write_title_block(ws, ds, opts, months, ly_months, extra_line=''):
    """匯出檔的標題區（A1-A6）。"""
    year = int(opts['year'])
    ws['A1'] = 'Shop YoY 分析報告'
    ws['A1'].font = Font(size=16, bold=True, color='0F172A')
    ws['A2'] = '資料檔：%s（上載時間：%s）' % (ds.get('filename', ''), ds.get('uploadedAt', ''))
    ws['A2'].font = Font(size=10, color='475569')
    ws['A3'] = '比較期間：%s  vs  %s' % (_period_label(months), _period_label(ly_months))
    ws['A3'].font = Font(size=11, bold=True, color='0F172A')

    levels = _resolve_levels(ds, opts)
    line = '分組：%s' % ' ＞ '.join(GROUP_LABELS.get(l, l) for l in levels)
    if extra_line:
        line += '　|　' + extra_line
    line += '　|　' + _filter_summary(opts)
    ws['A4'] = line
    ws['A4'].font = Font(size=10, color='475569')
    ws['A5'] = '產生時間：%s' % datetime.now().strftime('%Y-%m-%d %H:%M')
    ws['A5'].font = Font(size=10, color='475569')

    latest = ds.get('latestMonth', '')
    if latest and latest[:4] == str(year) and latest in months:
        ws['A6'] = '備註：%s 可能為不完整月份（若為進行中月份），YoY 數字請斟酌參考。' % latest
        ws['A6'].font = Font(size=10, bold=True, color='B45309')


def _row_metrics(g):
    """由組別資料計算 12 個顯示值：本期/去年/YoY% × (淨銷售、銷量、GP$)，再加 GP% 與 ±pp。"""
    sc, qc, gc = g['cur'][1], g['cur'][0], g['cur'][2]
    sl, ql, gl = g['ly'][1], g['ly'][0], g['ly'][2]
    pc = (gc / sc) if sc else None
    pl = (gl / sl) if sl else None
    y = lambda c, l: ((c - l) / l) if l and l > 0 else None
    return [
        sc, sl, y(sc, sl),
        qc, ql, y(qc, ql),
        gc, gl, y(gc, gl),
        pc, pl, (pc - pl) if (pc is not None and pl is not None) else None,
    ]


# ---------------------------------------------------------------------------
# Excel 匯出
# ---------------------------------------------------------------------------

def build_workbook(ds, opts, report):
    year = int(opts['year'])
    lyear = year - 1
    levels = _resolve_levels(ds, opts)
    months = sorted(str(m) for m in (opts.get('months') or []))
    ly_months = ['%d-%s' % (lyear, m[5:7]) for m in months]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'YoY'

    thin = Side(style='thin', color='D9E1EC')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_font = Font(bold=True, color='FFFFFF', size=11)
    head_fill = PatternFill('solid', fgColor='0F172A')
    sub_fill = PatternFill('solid', fgColor='F1F5F9')
    total_fill = PatternFill('solid', fgColor='EEF2FF')
    na_font = Font(color='94A3B8')
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left = Alignment(horizontal='left', vertical='center')
    right = Alignment(horizontal='right', vertical='center')

    FMT_QTY = '#,##0'
    FMT_MONEY = '#,##0'
    FMT_PCT = '0.00%'
    FMT_YOY = '+0.0%;-0.0%;0.0%'
    FMT_PP = '+0.00" pp";-0.00" pp";0.00" pp"'

    # ---- 標題區 ----
    _write_title_block(ws, ds, opts, months, ly_months)

    header_row = 8

    # ---- 欄位定義 ----
    metric_groups = [
        ('淨銷售', ['本期', '去年同期', 'YoY%']),
        ('銷量', ['本期', '去年同期', 'YoY%']),
        ('毛利 GP$', ['本期', '去年同期', 'YoY%']),
        ('毛利率', ['本期', '去年', '± pp']),
    ]
    if levels == ['shop']:
        lead = [
            ('店鋪', 'label', 26), ('Site', 'text', 9), ('店名', 'text', 18),
            ('Regional', 'text', 9), ('OM', 'text', 9), ('Class 1', 'text', 8),
            ('Class 2', 'text', 8), ('Size', 'text', 7), ('Type', 'text', 6), ('需關注', 'text', 7),
        ]
    else:
        lead_name = ' ＞ '.join(GROUP_LABELS.get(l, l) for l in levels) if len(levels) > 1 else GROUP_LABELS.get(levels[0], '分組')
        lead = [(lead_name, 'label', 30), ('店鋪數', 'qty', 9)]
    kinds = ['money', 'money', 'yoy', 'qty', 'qty', 'yoy', 'money', 'money', 'yoy', 'pct', 'pct', 'pp']
    kind_width = {'money': 15, 'qty': 11, 'yoy': 10, 'pct': 11, 'pp': 10, 'text': 9, 'label': 26}

    ncols = len(lead) + len(kinds)
    h1, h2 = header_row, header_row + 1

    # ---- 表頭 ----
    for rr in (h1, h2):
        for ci in range(1, ncols + 1):
            c = ws.cell(row=rr, column=ci)
            c.border = border
            c.alignment = center
            c.fill = head_fill if rr == h1 else sub_fill
            c.font = head_font if rr == h1 else Font(bold=True, color='334155')

    # 第一欄（及分組模式的店鋪數欄）合併兩行
    vmerged = [1] + ([2] if levels != ['shop'] else [])
    for ci in vmerged:
        ws.merge_cells(start_row=h1, start_column=ci, end_row=h2, end_column=ci)
        c = ws.cell(row=h1, column=ci)
        c.fill = head_fill
        c.font = head_font
        ws.cell(row=h2, column=ci).fill = head_fill

    ws.cell(row=h1, column=1, value=lead[0][0])
    col = 2
    if levels == ['shop']:
        ws.cell(row=h1, column=2, value='店鋪資料')
        ws.merge_cells(start_row=h1, start_column=2, end_row=h1, end_column=len(lead))
    else:
        ws.cell(row=h1, column=2, value='店鋪數')
    col = len(lead) + 1
    for gname, subs in metric_groups:
        ws.cell(row=h1, column=col, value=gname)
        if len(subs) > 1:
            ws.merge_cells(start_row=h1, start_column=col, end_row=h1, end_column=col + len(subs) - 1)
        col += len(subs)

    for ci, (name, kind, _w) in enumerate(lead, start=1):
        if ci in vmerged:
            continue
        ws.cell(row=h2, column=ci, value=name)
    col = len(lead) + 1
    for _gname, subs in metric_groups:
        for sname in subs:
            ws.cell(row=h2, column=col, value=sname)
            col += 1

    # 欄寬
    for ci, (name, kind, w) in enumerate(lead, start=1):
        ws.column_dimensions[get_column_letter(ci)].width = w
    col = len(lead) + 1
    for k in kinds:
        ws.column_dimensions[get_column_letter(col)].width = kind_width[k]
        col += 1

    # ---- 資料列 ----
    def write_cell(rr, ci, kind, v, bold=False, fill=None):
        c = ws.cell(row=rr, column=ci)
        c.border = border
        if fill:
            c.fill = fill
        if kind in ('label', 'text'):
            c.value = v if v is not None else ''
            c.font = Font(bold=bold)
            c.alignment = left if kind == 'label' else center
            return
        c.alignment = right
        if v is None:
            c.value = '—'
            c.font = na_font
            c.alignment = center
            return
        c.value = float(v)
        if kind == 'money':
            c.number_format = FMT_MONEY
        elif kind == 'qty':
            c.number_format = FMT_QTY
        elif kind == 'pct':
            c.number_format = FMT_PCT
        elif kind == 'yoy':
            c.number_format = FMT_YOY
        elif kind == 'pp':
            c.number_format = FMT_PP
        if kind in ('yoy', 'pp'):
            if v > 1e-9:
                c.font = Font(bold=bold, color='047857')
            elif v < -1e-9:
                c.font = Font(bold=bold, color='B91C1C')
            else:
                c.font = Font(bold=bold)
        else:
            c.font = Font(bold=bold)

    r = h2 + 1

    def emit(g, bold=False, fill=None):
        nonlocal r
        mvals = _row_metrics(g)
        if levels == ['shop']:
            lead_vals = [
                g['label'], g['code'], g['name'], g['regional'], g['om'],
                g['class1'], g['class2'], g['size'], g['type'],
                ('Y' if g.get('focus') else ''),
            ]
        else:
            lead_vals = [('　' * g.get('depth', 0)) + g['label'], g['n']]
        for ci, ((name, kind, _w), val) in enumerate(zip(lead, lead_vals), start=1):
            write_cell(r, ci, kind, val, bold=bold, fill=fill)
        for j, (kind, v) in enumerate(zip(kinds, mvals)):
            write_cell(r, len(lead) + 1 + j, kind, v, bold=bold, fill=fill)
        r += 1

    # 總計列（分組模式時「店鋪數」欄顯示通過篩選的店鋪總數）
    tot = report['totals']
    tot_g = {'cur': tot['cur'], 'ly': tot['ly'],
             'label': '總計', 'code': '', 'name': '', 'regional': '', 'om': '',
             'class1': '', 'class2': '', 'size': '', 'type': '', 'focus': False,
             'n': report.get('siteCount', 0)}
    emit(tot_g, bold=True, fill=total_fill)

    for g in report['groups']:
        emit(g)

    last_row = r - 1
    ws.auto_filter.ref = 'A%d:%s%d' % (h2, get_column_letter(ncols), last_row)
    ws.freeze_panes = 'B%d' % (h2 + 1)

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio


def build_matrix_workbook(ds, opts, matrix):
    """按月矩陣匯出：每組別一行、每個月一欄（顯示所選指標），尾加合計欄。"""
    year = int(opts['year'])
    lyear = year - 1
    levels = _resolve_levels(ds, opts)
    metric = opts.get('matrixMetric') or 'salesYoy'
    if metric == 'all':
        sections = [('salesYoy', '淨銷售 YoY%'), ('qtyYoy', '銷量 YoY%'),
                    ('gpYoy', '毛利 GP$ YoY%'), ('gpPp', '毛利率 ±pp')]
        metric_label = '全部指標（同步顯示）'
    elif metric == 'allv':
        sections = [('salesVal', '淨銷售（值）'), ('salesYoy', '淨銷售 YoY%'),
                    ('qtyVal', '銷量（值）'), ('qtyYoy', '銷量 YoY%'),
                    ('gpVal', '毛利 GP$（值）'), ('gpYoy', '毛利 GP$ YoY%'),
                    ('gpPctVal', '毛利率（值）'), ('gpPp', '毛利率 ±pp')]
        metric_label = '全部指標＋值（同步顯示）'
    else:
        _k, _label = MATRIX_METRICS.get(metric, ('yoy', '淨銷售 YoY%'))
        sections = [(metric, _label)]
        metric_label = _label
    multi = len(sections) > 1
    months = matrix['months']
    ly_months = ['%d-%s' % (lyear, m[5:7]) for m in months]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'YoY 按月'

    thin = Side(style='thin', color='D9E1EC')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_font = Font(bold=True, color='FFFFFF', size=11)
    head_fill = PatternFill('solid', fgColor='0F172A')
    total_fill = PatternFill('solid', fgColor='EEF2FF')
    sub_fill = PatternFill('solid', fgColor='F1F5F9')
    na_font = Font(color='94A3B8')
    border_sec = Border(left=Side(style='medium', color='94A3B8'), right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left = Alignment(horizontal='left', vertical='center')
    right = Alignment(horizontal='right', vertical='center')

    FMT = {'yoy': '+0.0%;-0.0%;0.0%', 'pp': '+0.00" pp";-0.00" pp";0.00" pp"',
           'money': '#,##0', 'qty': '#,##0', 'pct': '0.00%'}

    # ---- 標題區 ----
    _write_title_block(ws, ds, opts, months, ly_months,
                       extra_line='檢視：按月矩陣（指標：%s）' % metric_label)

    h1 = 8  # 表頭起始列
    h2 = h1 + 1 if multi else h1
    first_data = h2 + 1

    if levels == ['shop']:
        lead = [
            ('店鋪', 'label', 26), ('Site', 'text', 9), ('店名', 'text', 18),
            ('Regional', 'text', 9), ('OM', 'text', 9), ('Class 1', 'text', 8),
            ('Class 2', 'text', 8), ('Size', 'text', 7), ('Type', 'text', 6), ('需關注', 'text', 7),
        ]
        lead_group_name = '店鋪資料'
    else:
        lead_group_name = ' ＞ '.join(GROUP_LABELS.get(l, l) for l in levels) if len(levels) > 1 else GROUP_LABELS.get(levels[0], '分組')
        lead = [(lead_group_name, 'label', 30), ('店鋪數', 'qty', 9)]

    ncols = len(lead) + len(sections) * (len(months) + 1)

    # ---- 表頭 ----
    if multi:
        # 第一列：店鋪資料 + 各指標區塊（合併）
        for cc in range(1, len(lead) + 1):
            c = ws.cell(row=h1, column=cc)
            c.fill = head_fill
            c.border = border
        c = ws.cell(row=h1, column=1, value=lead_group_name)
        c.font = head_font
        c.alignment = center
        ws.merge_cells(start_row=h1, start_column=1, end_row=h1, end_column=len(lead))
        col = len(lead) + 1
        for _mk, mlabel in sections:
            n_sec = len(months) + 1
            for cc in range(col, col + n_sec):
                c = ws.cell(row=h1, column=cc)
                c.fill = head_fill
                c.border = border_sec if cc == col else border
            c = ws.cell(row=h1, column=col, value=mlabel)
            c.font = head_font
            c.alignment = center
            if n_sec > 1:
                ws.merge_cells(start_row=h1, start_column=col, end_row=h1, end_column=col + n_sec - 1)
            col += n_sec
        # 第二列：欄名 + 月份 + 合計
        for ci, (name, _k, _w) in enumerate(lead, start=1):
            c = ws.cell(row=h2, column=ci, value=name)
            c.border = border
            c.fill = sub_fill
            c.font = Font(bold=True, color='334155')
            c.alignment = center
        col = len(lead) + 1
        for _mk, _ml in sections:
            first = True
            for m in months:
                c = ws.cell(row=h2, column=col, value=m)
                c.border = border_sec if first else border
                c.fill = sub_fill
                c.font = Font(bold=True, color='334155')
                c.alignment = center
                first = False
                col += 1
            c = ws.cell(row=h2, column=col, value='合計')
            c.border = border
            c.fill = sub_fill
            c.font = Font(bold=True, color='334155')
            c.alignment = center
            col += 1
    else:
        # 單一指標：單列深色表頭
        for ci, (name, _k, _w) in enumerate(lead, start=1):
            c = ws.cell(row=h1, column=ci, value=name)
            c.border = border
            c.fill = head_fill
            c.font = head_font
            c.alignment = center
        col = len(lead) + 1
        for m in months:
            c = ws.cell(row=h1, column=col, value=m)
            c.border = border
            c.fill = head_fill
            c.font = head_font
            c.alignment = center
            col += 1
        c = ws.cell(row=h1, column=ncols, value='合計')
        c.border = border
        c.fill = head_fill
        c.font = head_font
        c.alignment = center

    # 欄寬
    for ci, (_n, _k, w) in enumerate(lead, start=1):
        ws.column_dimensions[get_column_letter(ci)].width = w
    for ci in range(len(lead) + 1, ncols + 1):
        ws.column_dimensions[get_column_letter(ci)].width = 11

    # ---- 儲存格寫入 ----
    def write_lead(rr, g, bold=False, fill=None):
        if levels == ['shop']:
            vals = [
                (g['label'], 'label'), (g['code'], 'text'), (g['name'], 'text'),
                (g['regional'], 'text'), (g['om'], 'text'), (g['class1'], 'text'),
                (g['class2'], 'text'), (g['size'], 'text'), (g['type'], 'text'),
                ('Y' if g.get('focus') else '', 'text'),
            ]
        else:
            vals = [(('　' * g.get('depth', 0)) + g['label'], 'label'), (g['n'], 'qty')]
        for ci, (v, k) in enumerate(vals, start=1):
            c = ws.cell(row=rr, column=ci)
            c.border = border
            if fill:
                c.fill = fill
            if k == 'qty':
                c.value = v if v is not None else ''
                c.number_format = FMT['qty']
                c.alignment = right
                c.font = Font(bold=bold)
            else:
                c.value = v if v is not None else ''
                c.alignment = left if k == 'label' else center
                c.font = Font(bold=bold)

    def write_cell(rr, ci, cur, ly, mk, bold=False, fill=None, sec_left=False):
        c = ws.cell(row=rr, column=ci)
        c.border = border_sec if sec_left else border
        c.alignment = right
        if fill:
            c.fill = fill
        v = _matrix_value(cur, ly, mk)
        if v is None:
            c.value = '—'
            c.font = na_font
            c.alignment = center
            return
        kind = MATRIX_METRICS.get(mk, ('yoy', ''))[0]
        c.value = float(v)
        c.number_format = FMT[kind]
        if kind in ('yoy', 'pp'):
            if v > 1e-9:
                c.font = Font(bold=bold, color='047857')
            elif v < -1e-9:
                c.font = Font(bold=bold, color='B91C1C')
            else:
                c.font = Font(bold=bold)
        else:
            c.font = Font(bold=bold)

    r = first_data

    # 總計列
    tot = matrix['totals']
    tot_g = {'label': '總計', 'n': matrix.get('siteCount', ''),
             'code': '', 'name': '', 'regional': '', 'om': '', 'class1': '',
             'class2': '', 'size': '', 'type': '', 'focus': False}
    write_lead(r, tot_g, bold=True, fill=total_fill)
    col = len(lead) + 1
    for mk, _ml in sections:
        first = True
        for m in months:
            write_cell(r, col, tot['cells'][m], tot['cellsLy'][m], mk, bold=True, fill=total_fill, sec_left=first)
            first = False
            col += 1
        write_cell(r, col, tot['aggCur'], tot['aggLy'], mk, bold=True, fill=total_fill)
        col += 1
    r += 1

    for g in matrix['groups']:
        write_lead(r, g)
        col = len(lead) + 1
        for mk, _ml in sections:
            first = True
            for m in months:
                write_cell(r, col, g['cells'][m], g['cellsLy'][m], mk, sec_left=first)
                first = False
                col += 1
            write_cell(r, col, g['aggCur'], g['aggLy'], mk)
            col += 1
        r += 1

    last_row = r - 1
    ws.auto_filter.ref = 'A%d:%s%d' % (h2, get_column_letter(ncols), last_row)
    ws.freeze_panes = 'B%d' % first_data

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio


# ---------------------------------------------------------------------------
# 路由
# ---------------------------------------------------------------------------

def _get_dataset(token):
    """取出未過期的 dataset；已過期即由記憶體刪除。"""
    ds = _DATASETS.get(token)
    if ds is None:
        return None
    if DATA_TTL_HOURS > 0 and time.time() - _TIMES.get(token, 0) > DATA_TTL_HOURS * 3600:
        with _LOCK:
            _DATASETS.pop(token, None)
            _TIMES.pop(token, None)
            if token in _DATASET_ORDER:
                _DATASET_ORDER.remove(token)
        return None
    return ds


@app.before_request
def _require_auth():
    """若設定了 APP_PASSWORD，所有請求（除 /healthz）需先通過 Basic Auth。"""
    if not APP_PASSWORD or request.path == '/healthz':
        return None
    auth = request.authorization
    if not auth or auth.password != APP_PASSWORD:
        return Response('需要登入：請輸入存取密碼（用戶名稱隨意）', 401,
                        {'WWW-Authenticate': 'Basic realm="Shop YoY"'})
    return None


@app.get('/')
def index():
    return render_template('index.html')


@app.get('/healthz')
def healthz():
    return jsonify({'ok': True})


def _store_dataset(ds):
    """存入記憶體（含 TTL 清理、最多保留數份），回傳 token。"""
    token = uuid.uuid4().hex
    now = time.time()
    with _LOCK:
        # 先清走過期資料，再存入新資料
        if DATA_TTL_HOURS > 0:
            cutoff = now - DATA_TTL_HOURS * 3600
            for t in [x for x in _DATASET_ORDER if _TIMES.get(x, 0) < cutoff]:
                _DATASET_ORDER.remove(t)
                _DATASETS.pop(t, None)
                _TIMES.pop(t, None)
        _DATASETS[token] = ds
        _TIMES[token] = now
        _DATASET_ORDER.append(token)
        while len(_DATASET_ORDER) > _DATASETS_MAX:
            old = _DATASET_ORDER.pop(0)
            _DATASETS.pop(old, None)
            _TIMES.pop(old, None)
    return token


def _cache_task_response(task_id):
    """完成後預先序列化回應（大檔自動 gzip）。

    輪詢 /api/progress 時可直接回傳同一份 bytes，避免每次輪詢都重新
    json.dumps＋gzip（大檔可省大量 CPU 同等候時間）；超過 8MB 就不暫存。
    """
    try:
        with _TASKS_LOCK:
            tk = _TASKS.get(task_id)
            if tk is None or tk.get('status') != 'done' or tk.get('ds') is None:
                return
            resp = {'ok': True, 'status': 'done', 'pct': 100.0, 'msg': '完成',
                    'token': tk.get('token'), 'dataset': tk['ds'],
                    'ttlHours': DATA_TTL_HOURS}
        payload = json.dumps(resp, ensure_ascii=False, separators=(',', ':')).encode('utf-8')
        if len(payload) > 8 * 1024 * 1024:
            return  # 太大就唔暫存，輪詢時即時序列化
        gz = gzip.compress(payload, 6) if len(payload) > 200 * 1024 else b''
        body, is_gz = (gz, True) if (gz and len(gz) < len(payload)) else (payload, False)
        with _TASKS_LOCK:
            tk = _TASKS.get(task_id)
            if tk is not None and tk.get('status') == 'done' and tk.get('ds') is not None:
                tk['resp_body'] = body
                tk['resp_gz'] = is_gz
                tk['resp_served'] = 0
    except Exception:  # noqa: BLE001 - 暫存失敗唔影響主流程
        pass


def _parse_task(task_id, filename, data):
    """背景解析任務：持續更新進度，完成後存入 _DATASETS。"""
    def progress(pct, msg=''):
        with _TASKS_LOCK:
            tk = _TASKS.get(task_id)
            if tk is not None:
                tk['pct'] = pct
                tk['msg'] = msg
                tk['ts'] = time.time()

    try:
        ds = parse_excel(data, progress=progress)
        ds['filename'] = filename
        ds['uploadedAt'] = datetime.now().strftime('%Y-%m-%d %H:%M')
        token = _store_dataset(ds)
        with _TASKS_LOCK:
            tk = _TASKS.get(task_id)
            if tk is not None:
                tk.update(status='done', pct=100.0, msg='完成', token=token, ds=ds, ts=time.time())
        _cache_task_response(task_id)
    except ValueError as e:
        with _TASKS_LOCK:
            tk = _TASKS.get(task_id)
            if tk is not None:
                tk.update(status='error', error=str(e), ts=time.time())
    except Exception as e:  # noqa: BLE001
        app.logger.exception('parse error')
        with _TASKS_LOCK:
            tk = _TASKS.get(task_id)
            if tk is not None:
                tk.update(status='error', error='解析檔案時發生錯誤：%s' % e, ts=time.time())


@app.post('/api/upload')
def api_upload():
    f = request.files.get('file')
    if f is None or not f.filename:
        return jsonify({'ok': False, 'error': '未收到檔案。'}), 400
    fn = f.filename
    if not fn.lower().endswith(('.xlsx', '.xlsm')):
        return jsonify({'ok': False, 'error': '只支援 .xlsx 檔案（請由 Hugo_by Shop_YoY 匯出）。'}), 400
    data = f.read()
    _prune_tasks()
    task_id = uuid.uuid4().hex
    with _TASKS_LOCK:
        _TASKS[task_id] = {'status': 'parsing', 'pct': 0.0, 'msg': '開始解析…', 'token': None,
                           'ds': None, 'error': None, 'ts': time.time()}
    threading.Thread(target=_parse_task, args=(task_id, fn, data), daemon=True).start()
    return jsonify({'ok': True, 'taskID': task_id})


@app.get('/api/progress/<task_id>')
def api_progress(task_id):
    """查詢解析進度；status = parsing / done / error。完成時一併回傳數據（大檔自動 gzip）。"""
    cached = None
    cached_gz = False
    resp = None
    accepts_gzip = 'gzip' in (request.headers.get('Accept-Encoding') or '')
    with _TASKS_LOCK:
        tk = _TASKS.get(task_id)
        if tk is None:
            return jsonify({'ok': False, 'error': '任務不存在或已過期，請重新上載。'}), 404
        body = tk.get('resp_body')
        if body is not None and (not tk.get('resp_gz') or accepts_gzip):
            cached, cached_gz = body, bool(tk.get('resp_gz'))
            tk['resp_served'] = tk.get('resp_served', 0) + 1
            if tk['resp_served'] >= 3:
                tk.pop('resp_body', None)  # 已送出數次；釋放暫存，之後即時序列化
        else:
            resp = {'ok': True, 'status': tk['status'], 'pct': round(tk['pct'], 1), 'msg': tk.get('msg') or ''}
            if tk['status'] == 'done':
                resp['token'] = tk['token']
                resp['dataset'] = tk['ds']
                resp['ttlHours'] = DATA_TTL_HOURS
            elif tk['status'] == 'error':
                resp['error'] = tk.get('error') or '解析失敗'

    headers = {'Content-Type': 'application/json; charset=utf-8', 'Cache-Control': 'no-store'}
    if cached is not None:
        if cached_gz:
            headers['Content-Encoding'] = 'gzip'
            headers['Vary'] = 'Accept-Encoding'
        return Response(cached, 200, headers)

    payload = json.dumps(resp, ensure_ascii=False, separators=(',', ':')).encode('utf-8')
    if len(payload) > 200 * 1024 and accepts_gzip:
        gz = gzip.compress(payload, 6)
        if len(gz) < len(payload):
            headers['Content-Encoding'] = 'gzip'
            headers['Vary'] = 'Accept-Encoding'
            return Response(gz, 200, headers)
    return Response(payload, 200, headers)


@app.post('/api/clear-all')
def api_clear_all():
    """即時清除所有已上載資料（由記憶體刪除）。"""
    with _LOCK:
        _DATASETS.clear()
        _DATASET_ORDER.clear()
        _TIMES.clear()
    return jsonify({'ok': True})


@app.post('/api/clear/<token>')
def api_clear(token):
    """即時清除指定的一份已上載資料。"""
    with _LOCK:
        _DATASETS.pop(token, None)
        _TIMES.pop(token, None)
        if token in _DATASET_ORDER:
            _DATASET_ORDER.remove(token)
    return jsonify({'ok': True})


@app.post('/api/report/<token>')
def api_report(token):
    ds = _get_dataset(token)
    if ds is None:
        return jsonify({'ok': False, 'error': '資料已過期，請重新上載 Excel。'}), 404
    opts = request.get_json(silent=True) or {}
    try:
        if (opts.get('view') or 'agg') == 'matrix':
            mx = compute_matrix(ds, opts)
            r2 = lambda arr: [round(x, 2) for x in arr]
            groups = [{
                'label': g['label'], 'n': g['n'], 'depth': g.get('depth', 0),
                'cells': {m: r2(g['cells'][m]) for m in mx['months']},
                'ly': {m: r2(g['cellsLy'][m]) for m in mx['months']},
                'agg': {'cur': r2(g['aggCur']), 'ly': r2(g['aggLy'])},
            } for g in mx['groups']]
            tot = mx['totals']
            return jsonify({
                'ok': True, 'view': 'matrix', 'months': mx['months'], 'groups': groups,
                'totals': {
                    'cells': {m: r2(tot['cells'][m]) for m in mx['months']},
                    'ly': {m: r2(tot['cellsLy'][m]) for m in mx['months']},
                    'agg': {'cur': r2(tot['aggCur']), 'ly': r2(tot['aggLy'])},
                },
                'siteCount': mx['siteCount'],
            })

        rep = compute_report(ds, opts)
    except Exception as e:  # noqa: BLE001
        return jsonify({'ok': False, 'error': '計算失敗：%s' % e}), 400
    groups = []
    for g in rep['groups']:
        groups.append({
            'label': g['label'], 'n': g['n'], 'depth': g.get('depth', 0),
            'cur': {'qty': g['cur'][0], 'sales': g['cur'][1], 'gp': g['cur'][2]},
            'ly': {'qty': g['ly'][0], 'sales': g['ly'][1], 'gp': g['ly'][2]},
        })
    tot = rep['totals']
    return jsonify({
        'ok': True,
        'groups': groups,
        'totals': {
            'cur': {'qty': tot['cur'][0], 'sales': tot['cur'][1], 'gp': tot['cur'][2]},
            'ly': {'qty': tot['ly'][0], 'sales': tot['ly'][1], 'gp': tot['ly'][2]},
        },
        'siteCount': rep['siteCount'],
    })


@app.post('/api/export/<token>')
def api_export(token):
    ds = _get_dataset(token)
    if ds is None:
        return jsonify({'ok': False, 'error': '資料已過期，請重新上載 Excel。'}), 404
    opts = request.get_json(silent=True) or {}
    if not opts.get('year') or not opts.get('months'):
        return jsonify({'ok': False, 'error': '請先選擇年份及至少一個月份。'}), 400
    try:
        if (opts.get('view') or 'agg') == 'matrix':
            mx = compute_matrix(ds, opts)
            bio = build_matrix_workbook(ds, opts, mx)
        else:
            rep = compute_report(ds, opts)
            bio = build_workbook(ds, opts, rep)
    except Exception as e:  # noqa: BLE001
        app.logger.exception('export error')
        return jsonify({'ok': False, 'error': '產生 Excel 時發生錯誤：%s' % e}), 500

    ys = sorted({int(str(m)[5:7]) for m in opts['months']})
    contiguous = all(ys[i] + 1 == ys[i + 1] for i in range(len(ys) - 1))
    if len(ys) == 1:
        mm = '%02d' % ys[0]
    elif contiguous:
        mm = '%02d-%02d' % (ys[0], ys[-1])
    else:
        mm = '-'.join('%02d' % m for m in ys)
    name = 'Shop_YoY_%d_%s.xlsx' % (int(opts['year']), mm)
    return send_file(
        bio, as_attachment=True, download_name=name,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@app.errorhandler(413)
def too_large(_e):
    return jsonify({'ok': False, 'error': '檔案過大（上限 32 MB）。'}), 413


if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5000))
    app.run(host='0.0.0.0', port=port, debug=False, threaded=True)
