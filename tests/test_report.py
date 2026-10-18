# -*- coding: utf-8 -*-
"""回歸測試：YoY 計算、篩選、排序、匯出排序。

重點保護「畫面與匯出用同一套排序」這條規則 —— 前端 index.html 的
compute() / metricsOf() / cellMetric() 是同一套邏輯的另一份實作，
改任何一邊都要同步更新這裡的期望值。
"""

import io

from openpyxl import load_workbook

import app as A

MONTHS = ['2025-01', '2025-02', '2026-01', '2026-02']

SITES = [
    dict(code='HA01', name='A 店', regional='R1', om='O1', class1='C1',
         class2='X', size='M', type='T1', focus=False, inTemplate=True, leanSku=False),
    dict(code='HA02', name='B 店', regional='R1', om='O1', class1='C1',
         class2='X', size='L', type='T1', focus=True, inTemplate=True, leanSku=False),
    dict(code='HA03', name='C 店', regional='R2', om='O2', class1='C2',
         class2='Y', size='S', type='T2', focus=False, inTemplate=False, leanSku=True),
    # 完全沒有資料的店鋪：用於驗證 hideEmpty
    dict(code='HA04', name='D 店', regional='R2', om='O2', class1='C2',
         class2='Y', size='S', type='T2', focus=False, inTemplate=False, leanSku=False),
]


def make_ds():
    """2026-01 vs 2025-01：HA01 +25%、HA02 +100%、HA03 無去年同期。"""
    rows = [
        [0, 0, 10, 80.0, 32.0],
        [0, 1, 5, 100.0, 25.0],
        [2, 0, 10, 100.0, 40.0],
        [2, 1, 5, 200.0, 50.0],
        [2, 2, 1, 50.0, 10.0],
        [1, 0, 10, 120.0, 48.0],
        [1, 1, 5, 100.0, 25.0],
        [1, 2, 2, 40.0, 8.0],
        [3, 0, 10, 150.0, 60.0],
        [3, 2, 2, 20.0, 4.0],
    ]
    return {
        'kind': 'shop',
        'months': MONTHS,
        'years': {'2025': ['2025-01', '2025-02'], '2026': ['2026-01', '2026-02']},
        'sites': [dict(s) for s in SITES],
        'rows': rows,
        'latestMonth': '2026-02',
        'rowCount': len(rows),
        'siteCount': len(SITES),
    }


def opts(**kw):
    base = dict(year=2026, months=['2026-01'], groupBys=['shop'], hideEmpty=True,
                sortKey='curSales', sortDir='desc')
    base.update(kw)
    return base


def labels(result):
    return [g['label'] for g in result['groups']]


def exported_codes(bio):
    """由匯出檔的資料列，按出現次序取出店鋪代碼。"""
    wb = load_workbook(io.BytesIO(bio.getvalue()))
    out = []
    for row in wb.active.iter_rows(values_only=True):
        for v in row:
            if isinstance(v, str) and v.startswith('HA0'):
                code = v[:4]
                if not out or out[-1] != code:
                    out.append(code)
                break
    return out


def test_yoy_and_totals():
    rep = A.compute_report(make_ds(), opts())
    by_code = {g['label'][:4]: g for g in rep['groups']}

    assert by_code['HA01']['cur'] == [10, 100.0, 40.0]
    assert by_code['HA01']['ly'] == [10, 80.0, 32.0]
    assert A._agg_sort_value(by_code['HA01'], 'yoySales') == 0.25

    assert A._agg_sort_value(by_code['HA02'], 'yoySales') == 1.0
    # 去年沒有資料 → 不計算 YoY
    assert A._agg_sort_value(by_code['HA03'], 'yoySales') is None

    assert rep['totals']['cur'] == [16, 350.0, 100.0]
    assert rep['totals']['ly'] == [15, 180.0, 57.0]


def test_sort_by_current_sales_desc_is_default():
    assert labels(A.compute_report(make_ds(), opts())) == ['HA02 B 店', 'HA01 A 店', 'HA03 C 店']


def test_sort_by_yoy_follows_screen():
    rep = A.compute_report(make_ds(), opts(sortKey='yoySales', sortDir='desc'))
    assert labels(rep) == ['HA02 B 店', 'HA01 A 店', 'HA03 C 店']

    rep = A.compute_report(make_ds(), opts(sortKey='yoySales', sortDir='asc'))
    assert labels(rep) == ['HA01 A 店', 'HA02 B 店', 'HA03 C 店']  # None 永遠排最後


def test_sort_by_label():
    rep = A.compute_report(make_ds(), opts(sortKey='label', sortDir='asc'))
    assert labels(rep) == ['HA01 A 店', 'HA02 B 店', 'HA03 C 店']

    rep = A.compute_report(make_ds(), opts(sortKey='label', sortDir='desc'))
    assert labels(rep) == ['HA03 C 店', 'HA02 B 店', 'HA01 A 店']


def test_unknown_sort_key_falls_back_to_sales():
    # 舊版前端沒傳 sortKey 時，行為要跟改動前一致
    assert labels(A.compute_report(make_ds(), opts(sortKey=None))) == \
        ['HA02 B 店', 'HA01 A 店', 'HA03 C 店']
    assert labels(A.compute_report(make_ds(), opts(sortKey='nonsense'))) == \
        ['HA02 B 店', 'HA01 A 店', 'HA03 C 店']


def test_hide_empty_site():
    rep = A.compute_report(make_ds(), opts())
    assert 'HA04 D 店' not in labels(rep)

    rep = A.compute_report(make_ds(), opts(hideEmpty=False))
    assert 'HA04 D 店' in labels(rep)


def test_site_and_attribute_filters():
    rep = A.compute_report(make_ds(), opts(siteCodes=['HA01', 'HA03']))
    assert labels(rep) == ['HA01 A 店', 'HA03 C 店']
    assert rep['siteCount'] == 2

    rep = A.compute_report(make_ds(), opts(focusOnly=True))
    assert labels(rep) == ['HA02 B 店']

    rep = A.compute_report(make_ds(), opts(templateOnly=True))
    assert labels(rep) == ['HA02 B 店', 'HA01 A 店']

    rep = A.compute_report(make_ds(), opts(regional=['R2']))
    assert labels(rep) == ['HA03 C 店']


def test_matrix_sort_by_month():
    rep = A.compute_matrix(make_ds(), opts(
        months=['2026-01', '2026-02'], sortKey='m:2026-02', sortDir='desc'))
    assert labels(rep) == ['HA01 A 店', 'HA03 C 店', 'HA02 B 店']


def test_matrix_sort_by_total():
    rep = A.compute_matrix(make_ds(), opts(
        months=['2026-01', '2026-02'], sortKey='total', matrixMetric='salesVal'))
    assert labels(rep) == ['HA01 A 店', 'HA02 B 店', 'HA03 C 店']

    ha01 = rep['groups'][0]
    assert ha01['aggCur'] == [20, 250.0, 100.0]
    assert ha01['aggLy'] == [20, 200.0, 80.0]


def test_export_row_order_matches_screen():
    ds = make_ds()
    for key, direction in (('curSales', 'desc'), ('yoySales', 'desc'), ('yoySales', 'asc')):
        o = opts(sortKey=key, sortDir=direction)
        rep = A.compute_report(ds, o)
        bio = A.build_workbook(ds, o, rep)
        assert exported_codes(bio) == [g['label'][:4] for g in rep['groups']], (key, direction)


def test_matrix_export_row_order_matches_screen():
    ds = make_ds()
    o = opts(months=['2026-01', '2026-02'], sortKey='m:2026-02',
             sortDir='desc', view='matrix', matrixMetric='salesYoy')
    mx = A.compute_matrix(ds, o)
    bio = A.build_matrix_workbook(ds, o, mx)
    assert exported_codes(bio) == [g['label'][:4] for g in mx['groups']]


def test_export_totals_match_report():
    ds = make_ds()
    o = opts()
    rep = A.compute_report(ds, o)
    wb = load_workbook(io.BytesIO(A.build_workbook(ds, o, rep).getvalue()))
    values = [c for row in wb.active.iter_rows(values_only=True) for c in row]

    # 匯出檔一定要包含本期合計淨銷售；用報表的合計值反查
    assert 350.0 in values
    assert 180.0 in values
